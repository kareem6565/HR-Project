"""Resume text extraction (Phase 4.4 / Prompt 5).

Extracts plain text from an uploaded resume file. Supports PDF via ``pypdf``
and DOCX via ``python-docx`` only. No OCR, no scanned-image support.

Validates:
- Real file signatures (%PDF-, PK) independently of file extensions.
- Encrypted or password-protected PDF detection.
- PDF page limit constraints.
- DOCX archive integrity: zip bomb uncompressed size limit, archive entry count limit,
  directory traversal checks, and required WordProcessingML structure ([Content_Types].xml,
  word/document.xml).
- Text extraction from both paragraphs and table cells.
- Controlled, structured error codes without leaking server filesystem paths or secrets.
"""

from __future__ import annotations

import io
import logging
import zipfile
from abc import ABC, abstractmethod
from pathlib import Path

from app.core.config import settings
from app.services.errors import (
    UPLOAD_ERROR_CORRUPT,
    UPLOAD_ERROR_ENCRYPTED,
    UPLOAD_ERROR_INVALID_ARCHIVE,
    UPLOAD_ERROR_NO_TEXT,
    UPLOAD_ERROR_PAGE_LIMIT_EXCEEDED,
    UPLOAD_ERROR_TOO_LARGE,
    UPLOAD_ERROR_UNSUPPORTED_TYPE,
    UploadValidationError,
    ValidationError,
)

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = frozenset({"pdf", "docx"})

# Content types (lower-cased) accepted alongside the matching extension.
PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

EXTENSION_TO_CONTENT_TYPE = {
    "pdf": PDF_CONTENT_TYPE,
    "docx": DOCX_CONTENT_TYPE,
}

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK"


class TextExtractor(ABC):
    """Extract plain text from one supported file format."""

    extension: str

    @abstractmethod
    def extract(self, data: bytes) -> str:
        """Return the extracted text, raising UploadValidationError if unextractable."""


class PdfTextExtractor(TextExtractor):
    extension = "pdf"

    def extract(self, data: bytes, *, max_pages: int | None = None) -> str:
        if not data:
            raise UploadValidationError("Uploaded PDF file is empty.", code=UPLOAD_ERROR_NO_TEXT)

        if not data.startswith(PDF_MAGIC):
            raise UploadValidationError(
                "File content does not match PDF signature.",
                code=UPLOAD_ERROR_UNSUPPORTED_TYPE,
            )

        from pypdf import PdfReader
        from pypdf.errors import PyPdfError

        effective_max_pages = max_pages if max_pages is not None else settings.max_pdf_pages

        try:
            reader = PdfReader(io.BytesIO(data))
        except (PyPdfError, ValueError, OSError, Exception) as exc:
            logger.warning("PDF reader initialization failed: %s", type(exc).__name__)
            raise UploadValidationError(
                "Uploaded PDF could not be read or is corrupt.",
                code=UPLOAD_ERROR_CORRUPT,
            ) from exc

        if reader.is_encrypted:
            raise UploadValidationError(
                "Encrypted or password-protected PDF files are not supported.",
                code=UPLOAD_ERROR_ENCRYPTED,
            )

        try:
            num_pages = len(reader.pages)
        except Exception as exc:
            raise UploadValidationError(
                "Failed to determine PDF page count.",
                code=UPLOAD_ERROR_CORRUPT,
            ) from exc

        if num_pages > effective_max_pages:
            raise UploadValidationError(
                f"PDF page count ({num_pages}) exceeds the maximum allowed limit of {effective_max_pages} pages.",
                code=UPLOAD_ERROR_PAGE_LIMIT_EXCEEDED,
            )

        pages: list[str] = []
        for idx, page in enumerate(reader.pages):
            try:
                t = page.extract_text() or ""
                pages.append(t)
            except Exception as exc:
                logger.warning("Error extracting text from PDF page %d: %s", idx, type(exc).__name__)
                raise UploadValidationError(
                    f"Failed to extract text from page {idx + 1} of PDF.",
                    code=UPLOAD_ERROR_CORRUPT,
                ) from exc

        text = "\n".join(pages).strip()
        if not text:
            raise UploadValidationError(
                "No extractable text was found in the uploaded file.",
                code=UPLOAD_ERROR_NO_TEXT,
            )
        return text


class DocxTextExtractor(TextExtractor):
    extension = "docx"

    def extract(
        self,
        data: bytes,
        *,
        max_uncompressed_bytes: int | None = None,
        max_entries: int | None = None,
    ) -> str:
        if not data:
            raise UploadValidationError("Uploaded DOCX file is empty.", code=UPLOAD_ERROR_NO_TEXT)

        if not data.startswith(ZIP_MAGIC):
            raise UploadValidationError(
                "File content does not match DOCX/ZIP signature.",
                code=UPLOAD_ERROR_UNSUPPORTED_TYPE,
            )

        from zipfile import BadZipFile, ZipFile

        from docx import Document

        effective_max_uncompressed = (
            max_uncompressed_bytes
            if max_uncompressed_bytes is not None
            else settings.max_docx_uncompressed_bytes
        )
        effective_max_entries = (
            max_entries if max_entries is not None else settings.max_docx_entries
        )

        # Inspect ZIP metadata in memory without filesystem extraction
        try:
            with ZipFile(io.BytesIO(data)) as zf:
                namelist = zf.namelist()
                if len(namelist) > effective_max_entries:
                    raise UploadValidationError(
                        f"DOCX archive entry count ({len(namelist)}) exceeds limit of {effective_max_entries}.",
                        code=UPLOAD_ERROR_INVALID_ARCHIVE,
                    )

                total_uncompressed = 0
                for info in zf.infolist():
                    if (
                        info.filename.startswith("/")
                        or info.filename.startswith("\\")
                        or ".." in info.filename
                    ):
                        raise UploadValidationError(
                            "DOCX archive contains unsafe file paths.",
                            code=UPLOAD_ERROR_INVALID_ARCHIVE,
                        )

                    total_uncompressed += info.file_size
                    if total_uncompressed > effective_max_uncompressed:
                        raise UploadValidationError(
                            f"DOCX declared uncompressed size exceeds maximum limit of {effective_max_uncompressed} bytes.",
                            code=UPLOAD_ERROR_TOO_LARGE,
                        )

                # Validate required WordProcessingML structural entries
                if "[Content_Types].xml" not in namelist or "word/document.xml" not in namelist:
                    raise UploadValidationError(
                        "Invalid DOCX archive: missing required WordProcessingML parts.",
                        code=UPLOAD_ERROR_INVALID_ARCHIVE,
                    )
        except BadZipFile as exc:
            raise UploadValidationError(
                "Uploaded DOCX archive is corrupt or invalid.",
                code=UPLOAD_ERROR_CORRUPT,
            ) from exc

        try:
            document = Document(io.BytesIO(data))
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]

            table_lines: list[str] = []
            for table in document.tables:
                for row in table.rows:
                    cell_texts = [
                        cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()
                    ]
                    if cell_texts:
                        table_lines.append(" | ".join(cell_texts))

            all_lines = paragraphs + table_lines
            text = "\n".join(all_lines).strip()
        except (BadZipFile, OSError, ValueError, Exception) as exc:
            logger.warning("DOCX document parsing failed: %s", type(exc).__name__)
            raise UploadValidationError(
                "Uploaded DOCX could not be read or is corrupt.",
                code=UPLOAD_ERROR_CORRUPT,
            ) from exc

        if not text:
            raise UploadValidationError(
                "No extractable text was found in the uploaded file.",
                code=UPLOAD_ERROR_NO_TEXT,
            )
        return text


_EXTRACTORS: dict[str, TextExtractor] = {
    PdfTextExtractor.extension: PdfTextExtractor(),
    DocxTextExtractor.extension: DocxTextExtractor(),
}


def extract_text(
    data: bytes,
    *,
    filename: str | None,
    max_pdf_pages: int | None = None,
    max_docx_uncompressed_bytes: int | None = None,
    max_docx_entries: int | None = None,
) -> str:
    """Extract text from resume bytes, validating type and extractability."""
    if not filename or not filename.strip():
        raise UploadValidationError(
            "An uploaded file must have a filename.",
            code=UPLOAD_ERROR_UNSUPPORTED_TYPE,
        )
    extension = Path(filename).suffix.lower().lstrip(".")
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise UploadValidationError(
            f"Unsupported file type: {extension or 'unknown'!r}. Only .pdf and .docx are allowed.",
            code=UPLOAD_ERROR_UNSUPPORTED_TYPE,
        )

    if isinstance(extractor, PdfTextExtractor):
        return extractor.extract(data, max_pages=max_pdf_pages)
    elif isinstance(extractor, DocxTextExtractor):
        return extractor.extract(
            data,
            max_uncompressed_bytes=max_docx_uncompressed_bytes,
            max_entries=max_docx_entries,
        )
    return extractor.extract(data)


def extension_for_filename(filename: str) -> str:
    """Return the lowercase extension (without dot), or raise if unsupported."""
    if not filename or not filename.strip():
        raise UploadValidationError(
            "An uploaded file must have a filename.",
            code=UPLOAD_ERROR_UNSUPPORTED_TYPE,
        )
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in SUPPORTED_EXTENSIONS:
        raise UploadValidationError(
            f"Unsupported file type: {extension or 'unknown'!r}. Only .pdf and .docx are allowed.",
            code=UPLOAD_ERROR_UNSUPPORTED_TYPE,
        )
    return extension
