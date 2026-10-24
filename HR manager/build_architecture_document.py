from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path(__file__).resolve().parent / "deliverables"
OUTPUT_FILE = OUTPUT_DIR / "BIC2_HR_Platform_Architecture_AR.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
MUTED = "5B6573"
WHITE = "FFFFFF"


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_cell_text(cell, text, bold=False, color=NAVY, size=10.3, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    set_rtl(paragraph, align)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def add_heading(doc, text, level=1):
    style = {1: "BIC2 Heading 1", 2: "BIC2 Heading 2", 3: "BIC2 Heading 3"}[level]
    paragraph = doc.add_paragraph(style=style)
    set_rtl(paragraph)
    paragraph.add_run(text)
    return paragraph


def add_body(doc, text, bold_label=None):
    paragraph = doc.add_paragraph(style="BIC2 Body")
    set_rtl(paragraph)
    if bold_label:
        label = paragraph.add_run(f"{bold_label}: ")
        set_run_font(label, size=10.8, color=NAVY, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.8, color=NAVY)
    return paragraph


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT)
    paragraph = cell.paragraphs[0]
    set_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(f"{title}: ")
    set_run_font(run, size=10.5, color=DARK_BLUE, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, widths)
    for cell, header in zip(table.rows[0].cells, headers):
        shade_cell(cell, LIGHT_BLUE)
        add_cell_text(cell, header, bold=True, color=DARK_BLUE, size=10.2)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_values):
            add_cell_text(cell, text, color=NAVY, size=9.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENTATION.PORTRAIT
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(NAVY)

    body = doc.styles.add_style("BIC2 Body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = normal
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("BIC2 Heading 1", 16, BLUE, 18, 10),
        ("BIC2 Heading 2", 13, BLUE, 14, 7),
        ("BIC2 Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    set_rtl(header)
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("BIC2 | توثيق معمارية المنصة")
    set_run_font(run, size=8.8, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    set_rtl(footer, WD_ALIGN_PARAGRAPH.CENTER)
    run = footer.add_run("منصة BIC2 لإدارة الموارد البشرية | صفحة ")
    set_run_font(run, size=8.5, color=MUTED)
    add_page_number(footer)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    set_rtl(title, WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_before = Pt(36)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("توثيق معمارية منصة BIC2")
    set_run_font(run, size=27, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    set_rtl(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
    subtitle.paragraph_format.space_after = Pt(26)
    run = subtitle.add_run("نظام الحضور والموارد البشرية المتكامل - النسخة الحالية MVP")
    set_run_font(run, size=13, color=DARK_BLUE)

    metadata = add_table(
        doc,
        ["البند", "القيمة"],
        [
            ("نوع المستند", "مرجع معماري ووظيفي للنظام الحالي"),
            ("بيئة التشغيل", "XAMPP محليًا: Apache + PHP + MySQL"),
            ("قاعدة البيانات", "attendance_db عبر PDO وMySQL"),
            ("واجهة الاستخدام", "Arabic RTL - HTML/CSS/Tailwind + JavaScript"),
            ("تاريخ التوثيق", "03 سبتمبر 2026"),
        ],
        [2700, 6660],
    )
    metadata.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    add_callout(
        doc,
        "ملخص",
        "BIC2 منصة داخلية لإدارة المستخدمين والحضور الجغرافي والإجازات والرواتب والوثائق والإشعارات. "
        "تدعم أدوار Admin وHR وManager وEmployee مع صلاحيات ومسارات عمل مختلفة.",
    )

    add_heading(doc, "1. نطاق النظام الحالي", 1)
    add_body(doc, "يوحّد النظام إدارة الحضور اليومي والخدمات الذاتية للموظف وعمليات الموارد البشرية في تطبيق ويب واحد يعمل محليًا عبر XAMPP.")
    add_table(
        doc,
        ["المجال", "الوظائف المتاحة الآن"],
        [
            ("الحسابات والصلاحيات", "تسجيل دخول، جلسات آمنة، توجيه حسب الدور، ومنع الوصول غير المصرح به."),
            ("الحضور", "تسجيل حضور وانصراف، شيفتات نهارية وليلية، استراحة، سجل آخر 30 يومًا."),
            ("الموقع الجغرافي", "نطاق شركة/فئة، تحقق عند الحضور والانصراف والاستراحة، ومراقبة أثناء الدوام."),
            ("HR", "إجازات، رواتب، وثائق، خطابات، إشعارات، وتقارير قابلة للتصدير."),
            ("الإدارة", "إدارة الموظفين والفئات والشيفتات والموقع وإسناد الموظفين ومتابعة الحضور."),
            ("إنجازات العمل", "رفع ملفات وروابط وتقارير الموظف ومراجعتها إداريًا."),
        ],
        [2700, 6660],
    )

    add_heading(doc, "2. المعمارية التقنية", 1)
    add_body(doc, "يعتمد النظام على معمارية طبقات بسيطة. تتعامل صفحات الواجهة مع واجهات API بصيغة JSON، وتتولى طبقة PHP التحقق من الجلسة والصلاحية ومنطق العمل، ثم تحفظ البيانات في MySQL.")
    add_table(
        doc,
        ["الطبقة", "المكونات والمسؤولية"],
        [
            ("واجهة المستخدم", "login.php، employee.php، admin.php، hr_suite.php، profile.php، work.php. واجهات عربية RTL وتفاعلات JavaScript."),
            ("واجهة الخدمات", "api/attendance.php، api/employees.php، api/categories.php، api/settings.php، api/hr.php، api/work.php."),
            ("منطق الأعمال", "includes/auth.php للجلسات والأدوار، includes/helpers.php للوقت والموقع، includes/hr_platform.php للإشعارات وسير عمل HR."),
            ("البيانات", "MySQL: users وattendance وcategories وcompany_settings ووحدات HR وسجل geofence_events."),
            ("التخزين", "مجلد uploads/hr_documents لحفظ الوثائق والمرفقات المصرح بها."),
        ],
        [2700, 6660],
    )
    add_callout(doc, "مسار البيانات", "المستخدم > صفحة الواجهة > API بصيغة JSON > فحص الجلسة والصلاحية > منطق الأعمال > MySQL > استجابة الواجهة.")

    add_heading(doc, "3. الأدوار والصلاحيات", 1)
    add_table(
        doc,
        ["الدور", "النطاق الوظيفي"],
        [
            ("Admin", "إدارة كاملة للموظفين والفئات والإعدادات والحضور، والوصول إلى بوابة HR والتقارير."),
            ("HR", "إدارة الموظفين والفئات وإعدادات موقع الشركة، متابعة الحضور، الإجازات والرواتب والوثائق والخطابات والإشعارات."),
            ("Manager", "عرض فريقه ومتابعة طلبات الفريق بحسب منطق بوابة HR."),
            ("Employee", "تسجيل الحضور والاستراحة والانصراف، طلب الإجازة، متابعة الكشوف، الوثائق والخطابات، ورفع الإنجازات."),
        ],
        [2100, 7260],
    )
    add_callout(doc, "ملاحظة تشغيلية", "تستخدم جميع تبويبات المتصفح العادي نفس جلسة الدخول. لاختبار حسابين في وقت واحد، استخدم نافذة خاصة أو متصفحًا مختلفًا.")

    add_heading(doc, "4. الحضور، الشيفتات، والاستراحة", 1)
    add_body(doc, "يوفر api/attendance.php منطقًا موحدًا للحضور والانصراف والاستراحة، مع دعم الشيفتات التي تتخطى منتصف الليل.")
    add_table(
        doc,
        ["التدفق", "الضوابط الحالية"],
        [
            ("تسجيل الحضور", "يتطلب إحداثيات صحيحة ووجود المستخدم داخل نطاق الفئة/الشركة. يسجل موقع الدخول والمسافة والحالة."),
            ("تسجيل الانصراف", "يتطلب موقعًا داخل النطاق ويسجل موقع الخروج والمسافة. يغلق الاستراحة المفتوحة ويحسب تجاوزها."),
            ("بدء/إنهاء الاستراحة", "يتطلب حضورًا مفتوحًا وموقعًا داخل النطاق. يحسب مدة الاستراحة والتجاوز بالدقائق."),
            ("المراقبة الحية", "بوابة الموظف ترسل location_ping عبر geolocation أثناء وجود الصفحة مفتوحة؛ تسجل الخروج والعودة للنطاق."),
            ("الشيفت الليلي", "يقرأ السجل المفتوح من اليوم أو اليوم السابق لمعالجة الحضور الذي يتجاوز منتصف الليل."),
        ],
        [2700, 6660],
    )

    add_heading(doc, "5. النطاق الجغرافي والتنبيهات", 1)
    add_body(doc, "يُحدد Admin أو HR إحداثيات الشركة ونصف قطرها من إعدادات الشركة. عند الحفظ، تُحدّث الفئات لتستخدم الإعداد نفسه. ويمكن لكل فئة حفظ إعداد جغرافي مستقل عند الحاجة.")
    add_table(
        doc,
        ["العنصر", "التفاصيل"],
        [
            ("مصدر الموقع", "Browser Geolocation من جهاز المستخدم بعد منحه الإذن."),
            ("التحقق", "حساب Haversine للمسافة بين موقع المستخدم ومركز الفئة، ثم مقارنة النتيجة بنطاق الأمتار المسموح."),
            ("الأحداث المسجلة", "checkin_inside، checkout_inside، break_started_inside، break_ended_inside، left_geofence، returned_to_geofence."),
            ("التنبيه", "عند الخروج من النطاق تُنشأ إشعارات لمستخدمي Admin وHR، ويُحفظ الحدث في geofence_events."),
            ("القيود", "تطبيق ويب يمكنه الرصد والتنبيه ولا يمكنه منع الشخص جسديًا من مغادرة الموقع. يعمل التتبع ما دامت الصفحة مفتوحة."),
        ],
        [2700, 6660],
    )

    add_heading(doc, "6. وحدات الموارد البشرية", 1)
    add_table(
        doc,
        ["الوحدة", "المكونات الرئيسية"],
        [
            ("الإجازات", "أنواع إجازة، أرصدة سنوية، طلبات، حالة موافقات، ومتابعة للموظف والمدير وHR."),
            ("الرواتب", "دورات راتب شهرية، بدلات وخصومات، عناصر كشوف، وإتاحة كشف الدفع."),
            ("الوثائق والخطابات", "رفع وثائق الموظف، تحديد صلاحيتها، طلب وإصدار خطابات HR."),
            ("الإشعارات", "إشعارات داخلية لقرارات الإجازات والرواتب والأحداث الجغرافية."),
            ("التقارير", "مصادر تقارير قابلة للتصدير إلى Excel/PDF حسب وحدات بوابة HR."),
        ],
        [2700, 6660],
    )

    add_heading(doc, "7. نموذج البيانات", 1)
    add_table(
        doc,
        ["المجموعة", "الجداول الأساسية"],
        [
            ("المستخدمون والهيكل", "users، categories، company_settings."),
            ("الحضور والموقع", "attendance، geofence_events."),
            ("الإجازات", "leave_types، leave_balances، leave_requests."),
            ("الرواتب", "payroll_runs، payroll_adjustments، payroll_items."),
            ("ملفات HR", "employee_documents، hr_letters، notifications، audit_logs."),
            ("العمل", "work_submissions وما يرتبط بها من مرفقات وروابط وحالة مراجعة."),
        ],
        [2700, 6660],
    )
    add_body(doc, "يتعامل PDO مع قاعدة البيانات باستعلامات مجهزة، وتستخدم مفاتيح خارجية في الجداول المهمة لضمان اتساق المستخدمين وسجلاتهم.")

    add_heading(doc, "8. التشغيل المحلي", 1)
    add_table(
        doc,
        ["الخطوة", "الإجراء"],
        [
            ("1", "تشغيل Apache وMySQL من XAMPP."),
            ("2", "فتح http://localhost/BIC2_Mvp/login.php."),
            ("3", "التأكد من قاعدة attendance_db وإعدادات الاتصال في includes/db.php دون وضع كلمة مرور في المستودع."),
            ("4", "تسجيل الدخول بالدور المطلوب. حسابات الاختبار يجب تغيير كلمات مرورها قبل التسليم الفعلي."),
            ("5", "قبل تجربة الحضور، ضبط موقع الشركة/الفئة ليطابق الموقع الفعلي ومنح المتصفح صلاحية الموقع."),
        ],
        [1200, 8160],
    )

    add_heading(doc, "9. إعدادات مؤقتة للتجربة", 1)
    add_callout(
        doc,
        "تنبيه",
        "تم تعطيل مهلة الحضور مؤقتًا داخل api/attendance.php عبر ENFORCE_CHECKIN_CUTOFF = false لتسهيل تجربة النظام خارج وقت الشفت. "
        "يجب إعادتها إلى true قبل التشغيل الفعلي.",
    )
    add_body(doc, "تظل شروط الموقع الجغرافي، وتوقيت الانصراف، وقواعد الاستراحة فعالة ما لم تُغيّر صراحة. كما يجب مراجعة إحداثيات الشركة ونطاقها قبل الاختبار؛ الإحداثيات الافتراضية في المشروع كانت للرياض ونطاق 300 متر.")

    add_heading(doc, "10. نقاط تحسين مقترحة قبل الإنتاج", 1)
    add_table(
        doc,
        ["الأولوية", "التحسين المقترح"],
        [
            ("مرتفعة", "نقل بيانات الاتصال وكلمات المرور إلى متغيرات بيئة، وتغيير حسابات التجربة، وتشغيل HTTPS."),
            ("مرتفعة", "تطبيق سياسة نسخة احتياطية لقاعدة البيانات والملفات وتهيئة سجل مراقبة مركزي."),
            ("متوسطة", "تطبيق موبايل أو MDM للتتبع في الخلفية والتحقق الأقوى من تزوير الموقع عند الحاجة."),
            ("متوسطة", "إضافة شاشة مرئية لأحداث geofence_events داخل لوحة HR مع فلاتر تاريخ وموظف."),
            ("متوسطة", "تفعيل البريد أو واتساب عبر مزود رسمي بعد تحديد المتطلبات وسياسة الخصوصية."),
        ],
        [1800, 7560],
    )

    add_callout(doc, "خلاصة", "النسخة الحالية صالحة كـMVP متكامل للتجربة المحلية. قبل النشر الفعلي، يجب تثبيت سياسات الأمن والتشغيل والنسخ الاحتياطي وتحديد متطلبات التتبع الجغرافي بدقة.")

    OUTPUT_DIR.mkdir(exist_ok=True)
    doc.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    build_document()
